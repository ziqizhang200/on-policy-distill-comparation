#!/usr/bin/env python3
"""按导师意见重构开题报告，并保留原有 HIT 模板样式。"""

from __future__ import annotations

import argparse
import os
import re
from copy import deepcopy
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches


DEFAULT_REPORT = Path(
    "Papers/output/开题报告-基于在线策略蒸馏的大语言模型后训练方法研究(1).docx"
)


REFERENCE_ENTRIES = [
    "Hinton G, Vinyals O, Dean J. Distilling the Knowledge in a Neural Network[EB/OL]. arXiv:1503.02531, 2015.",
    "Kim Y, Rush A M. Sequence-Level Knowledge Distillation[C]//Proceedings of the 2016 Conference on Empirical Methods in Natural Language Processing. ACL, 2016: 1317-1327. DOI:10.18653/v1/D16-1139.",
    "Ross S, Gordon G J, Bagnell J A. A Reduction of Imitation Learning and Structured Prediction to No-Regret Online Learning[C]//Proceedings of the Fourteenth International Conference on Artificial Intelligence and Statistics. PMLR, 2011, 15: 627-635.",
    "Lin A, Wohlwend J, Chen H, et al. Autoregressive Knowledge Distillation through Imitation Learning[C]//Proceedings of the 2020 Conference on Empirical Methods in Natural Language Processing. ACL, 2020: 6121-6133. DOI:10.18653/v1/2020.emnlp-main.494.",
    "Wen Y, Li Z, Du W, et al. f-Divergence Minimization for Sequence-Level Knowledge Distillation[C]//Proceedings of the 61st Annual Meeting of the Association for Computational Linguistics. ACL, 2023: 10817-10834. DOI:10.18653/v1/2023.acl-long.605.",
    "Agarwal R, Vieillard N, Zhou Y, et al. On-Policy Distillation of Language Models: Learning from Self-Generated Mistakes[C]//International Conference on Learning Representations. 2024. arXiv:2306.13649.",
    "Gu Y, Dong L, Wei F, et al. MiniLLM: Knowledge Distillation of Large Language Models[C]//International Conference on Learning Representations. 2024. arXiv:2306.08543.",
    "Ko J, Kim S, Chen T, et al. DistiLLM: Towards Streamlined Distillation for Large Language Models[C]//Proceedings of the 41st International Conference on Machine Learning. PMLR, 2024, 235: 24872-24895.",
    "Ko J, Chen T, Kim S, et al. DistiLLM-2: A Contrastive Approach Boosts the Distillation of LLMs[C]//Proceedings of the 42nd International Conference on Machine Learning. PMLR, 2025, 267: 31044-31062.",
    "Li Y, Zuo Y, He B, et al. Rethinking On-Policy Distillation of Large Language Models: Phenomenology, Mechanism, and Recipe[EB/OL]. arXiv:2604.13016v2, 2026. DOI:10.48550/arXiv.2604.13016.",
    "Zhu W, Xie R, Wang R, et al. Hybrid Policy Distillation for LLMs[C]//International Conference on Machine Learning. 2026. arXiv:2604.20244v2.",
    "Qwen Team. Qwen3 Technical Report[EB/OL]. arXiv:2505.09388, 2025. DOI:10.48550/arXiv.2505.09388.",
    "Kullback S, Leibler R A. On Information and Sufficiency[J]. The Annals of Mathematical Statistics, 1951, 22(1): 79-86. DOI:10.1214/aoms/1177729694.",
    "Lin J. Divergence Measures Based on the Shannon Entropy[J]. IEEE Transactions on Information Theory, 1991, 37(1): 145-151.",
    "Hu E J, Shen Y, Wallis P, et al. LoRA: Low-Rank Adaptation of Large Language Models[C]//International Conference on Learning Representations. 2022. arXiv:2106.09685.",
    "Kwon W, Li Z, Zhuang S, et al. Efficient Memory Management for Large Language Model Serving with PagedAttention[C]//Proceedings of the 29th Symposium on Operating Systems Principles. ACM, 2023: 611-626. DOI:10.1145/3600006.3613165.",
    "Yao S, Zhao J, Yu D, et al. ReAct: Synergizing Reasoning and Acting in Language Models[C]//International Conference on Learning Representations. 2023. arXiv:2210.03629.",
    "Li M, Zhao Y, Yu B, et al. API-Bank: A Comprehensive Benchmark for Tool-Augmented LLMs[C]//Proceedings of the 2023 Conference on Empirical Methods in Natural Language Processing. ACL, 2023: 3102-3116. DOI:10.18653/v1/2023.emnlp-main.187.",
    "Liu Z, Hoang T, Zhang J, et al. APIGen: Automated Pipeline for Generating Verifiable and Diverse Function-Calling Datasets[C]//Advances in Neural Information Processing Systems. 2024. DOI:10.52202/079017-1725.",
    "Liu W, Huang X, Zeng X, et al. ToolACE: Winning the Points of LLM Function Calling[C]//International Conference on Learning Representations. 2025. arXiv:2409.00920.",
    "Patil S G, Mao H, Yan F, et al. The Berkeley Function Calling Leaderboard (BFCL): From Tool Use to Agentic Evaluation of Large Language Models[C]//Proceedings of the 42nd International Conference on Machine Learning. PMLR, 2025, 267: 48371-48392.",
]


def norm(text: str) -> str:
    return " ".join(text.split())


def find_paragraph(doc: Document, exact: str):
    for paragraph in doc.paragraphs:
        if norm(paragraph.text) == norm(exact):
            return paragraph
    raise ValueError(f"未找到段落：{exact}")


def clear_runs_keep_structure(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag in {qn("w:r"), qn("w:hyperlink"), qn("w:smartTag")}:
            paragraph._p.remove(child)


def set_paragraph_text(paragraph, text: str) -> None:
    clear_runs_keep_structure(paragraph)
    paragraph.add_run(text)


def set_plain_paragraph_text(paragraph, text: str) -> None:
    """将含自动编号域的段落改为稳定的静态文本，同时保留段落格式。"""
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.add_run(text)


def remove_between(start_paragraph, end_paragraph=None) -> None:
    node = start_paragraph._p.getnext()
    end_node = end_paragraph._p if end_paragraph is not None else None
    while node is not None and node is not end_node:
        nxt = node.getnext()
        if node.tag != qn("w:sectPr"):
            node.getparent().remove(node)
        node = nxt


def detach(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def replace_ppr(paragraph, ppr_template) -> None:
    current = paragraph._p.pPr
    if current is not None:
        paragraph._p.remove(current)
    if ppr_template is not None:
        paragraph._p.insert(0, deepcopy(ppr_template))


def insert_paragraph_before(doc, anchor, text, ppr_template=None):
    paragraph = doc.add_paragraph()
    replace_ppr(paragraph, ppr_template)
    paragraph.add_run(text)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def set_keep_with_next(paragraph, keep: bool) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    keep_next = ppr.find(qn("w:keepNext"))
    if keep and keep_next is None:
        ppr.append(OxmlElement("w:keepNext"))
    elif not keep and keep_next is not None:
        ppr.remove(keep_next)


def insert_element_before(anchor, element) -> None:
    anchor._p.addprevious(element)


def make_math_paragraph(nodes):
    paragraph = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:firstLine"), "0")
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    ppr.extend([ind, jc])
    paragraph.append(ppr)
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    for node in nodes:
        math.append(node)
    math_para.append(math)
    paragraph.append(math_para)
    return paragraph


def m_run(text: str):
    run = OxmlElement("m:r")
    run_pr = OxmlElement("m:rPr")
    run.append(run_pr)
    token = OxmlElement("m:t")
    token.text = text
    run.append(token)
    return run


def m_sub(base: str, subscript: str):
    element = OxmlElement("m:sSub")
    element.append(OxmlElement("m:sSubPr"))
    base_el = OxmlElement("m:e")
    base_el.append(m_run(base))
    sub_el = OxmlElement("m:sub")
    sub_el.append(m_run(subscript))
    element.extend([base_el, sub_el])
    return element


def m_subsup(base: str, subscript: str, superscript: str):
    element = OxmlElement("m:sSubSup")
    element.append(OxmlElement("m:sSubSupPr"))
    base_el = OxmlElement("m:e")
    base_el.append(m_run(base))
    sub_el = OxmlElement("m:sub")
    sub_el.append(m_run(subscript))
    sup_el = OxmlElement("m:sup")
    sup_el.append(m_run(superscript))
    element.extend([base_el, sub_el, sup_el])
    return element


def m_fraction(numerator_nodes, denominator_nodes):
    element = OxmlElement("m:f")
    element.append(OxmlElement("m:fPr"))
    num = OxmlElement("m:num")
    den = OxmlElement("m:den")
    for node in numerator_nodes:
        num.append(node)
    for node in denominator_nodes:
        den.append(node)
    element.extend([num, den])
    return element


def m_delimiter(nodes):
    element = OxmlElement("m:d")
    element.append(OxmlElement("m:dPr"))
    expr = OxmlElement("m:e")
    for node in nodes:
        expr.append(node)
    element.append(expr)
    return element


def rollout_equation():
    return make_math_paragraph(
        [
            m_sub("y", "1:T"),
            m_run(" ∼ "),
            m_sub("π", "θ"),
            m_delimiter([m_run("· | x")]),
            m_run(",    "),
            m_sub("s", "t"),
            m_run(" = "),
            m_delimiter([m_run("x, "), m_sub("y", "<t")]),
        ]
    )


def distribution_equation():
    return make_math_paragraph(
        [
            m_sub("p", "S,t"),
            m_run(" = "),
            m_sub("π", "θ"),
            m_delimiter([m_run("· | "), m_sub("s", "t")]),
            m_run(",    "),
            m_sub("p", "T,t"),
            m_run(" = "),
            m_sub("π", "φ"),
            m_delimiter([m_run("· | "), m_sub("s", "t")]),
        ]
    )


def objective_equation():
    numerator = [
        m_subsup("∑", "t=1", "T"),
        m_sub("m", "t"),
        m_sub("D", "δ"),
        m_delimiter([m_sub("p", "T,t"), m_run(" ∥ "), m_sub("p", "S,t")]),
    ]
    denominator = [m_subsup("∑", "t=1", "T"), m_sub("m", "t")]
    return make_math_paragraph(
        [
            m_run("θ*"),
            m_run(" = "),
            m_sub("arg min", "θ"),
            m_sub("E", "x∼𝒟, y∼πθ"),
            m_delimiter([m_fraction(numerator, denominator)]),
        ]
    )


def topk_equation():
    return make_math_paragraph(
        [
            m_subsup("S", "t", "(K)"),
            m_run(" = TopK"),
            m_delimiter([m_sub("p", "T,t")]),
            m_run(",    "),
            m_subsup("Z", "T,t", "(K)"),
            m_run(" = "),
            m_sub("∑", "v∈Sₜ⁽ᴷ⁾"),
            m_sub("p", "T,t"),
            m_delimiter([m_run("v")]),
        ]
    )


def topk_normalized_equation():
    return make_math_paragraph(
        [
            m_subsup("p̃", "T,t", "(K)"),
            m_delimiter([m_run("v")]),
            m_run(" = "),
            m_fraction(
                [m_sub("p", "T,t"), m_delimiter([m_run("v")])],
                [m_subsup("Z", "T,t", "(K)")],
            ),
            m_run(",   v ∈ "),
            m_subsup("S", "t", "(K)"),
        ]
    )


def k1_equation():
    return make_math_paragraph(
        [
            m_sub("v", "t"),
            m_run(" ∼ "),
            m_sub("q", "t"),
            m_run(",    "),
            m_sub("Ê", "pT,t"),
            m_delimiter([m_run("g(v)")]),
            m_run(" = "),
            m_fraction(
                [m_sub("p", "T,t"), m_delimiter([m_sub("v", "t")])],
                [m_sub("q", "t"), m_delimiter([m_sub("v", "t")])],
            ),
            m_run("g"),
            m_delimiter([m_sub("v", "t")]),
        ]
    )


def wrap_text(draw, text, font, max_width):
    lines = []
    current = ""
    for ch in text:
        candidate = current + ch
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return "\n".join(lines)


def draw_centered_text(draw, box, text, font, fill="#111827", spacing=12):
    x0, y0, x1, y1 = box
    wrapped = wrap_text(draw, text, font, x1 - x0 - 70)
    bbox = draw.multiline_textbbox((0, 0), wrapped, font=font, spacing=spacing, align="center")
    x = (x0 + x1 - (bbox[2] - bbox[0])) / 2
    y = (y0 + y1 - (bbox[3] - bbox[1])) / 2
    draw.multiline_text((x, y), wrapped, font=font, fill=fill, spacing=spacing, align="center")


def arrow(draw, start, end, fill="#475569", width=8):
    draw.line([start, end], fill=fill, width=width)
    x, y = end
    draw.polygon([(x, y), (x - 18, y - 30), (x + 18, y - 30)], fill=fill)


def create_relationship_diagram(path: Path) -> None:
    width, height = 2400, 1500
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    heading_font = ImageFont.truetype("/System/Library/Fonts/STHeiti Medium.ttc", 52)
    layer_font = ImageFont.truetype("/System/Library/Fonts/STHeiti Medium.ttc", 38)
    body_font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Songti.ttc", 34)
    small_font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Songti.ttc", 29)

    draw.text((120, 58), "在线策略蒸馏的研究问题分解与闭环关系", font=heading_font, fill="#0F172A")

    state_box = (250, 165, 2150, 405)
    objective_box = (250, 500, 2150, 740)
    estimator_box = (250, 835, 2150, 1075)
    outcome_box = (250, 1195, 2150, 1435)
    boxes = [
        (
            state_box,
            "#E8F1FB",
            "状态分布层（研究问题一）\n学生 rollout → 实际生成前缀 → 教师同状态反馈\n决定教师监督作用于何种状态分布",
        ),
        (
            objective_box,
            "#EAF7EF",
            "目标函数层（研究问题二）\nForward KL · Reverse KL · 广义散度\n散度方向决定概率质量的覆盖、集中与梯度行为",
        ),
        (
            estimator_box,
            "#FFF6DF",
            "信息估计层（研究问题三）\n全词表 · Top-K · K1 sampled-token\n教师信息预算决定估计偏差、方差与计算开销",
        ),
        (
            outcome_box,
            "#F2ECFA",
            "联合决定学生参数更新\n优化稳定性  ·  Function Calling 正确性  ·  教师反馈成本",
        ),
    ]
    for box, color, text in boxes:
        draw.rounded_rectangle(box, radius=26, fill=color, outline="#64748B", width=4)
        first, *rest = text.split("\n")
        first_bbox = draw.textbbox((0, 0), first, font=layer_font)
        first_x = (box[0] + box[2] - (first_bbox[2] - first_bbox[0])) / 2
        draw.text((first_x, box[1] + 30), first, font=layer_font, fill="#0F172A")
        rest_text = "\n".join(rest)
        rest_bbox = draw.multiline_textbbox((0, 0), rest_text, font=body_font, spacing=12, align="center")
        rest_x = (box[0] + box[2] - (rest_bbox[2] - rest_bbox[0])) / 2
        draw.multiline_text((rest_x, box[1] + 100), rest_text, font=body_font, fill="#334155", spacing=12, align="center")

    arrow(draw, (1200, 405), (1200, 500))
    draw.text((1260, 430), "在同一状态上比较师生分布", font=small_font, fill="#475569")
    arrow(draw, (1200, 740), (1200, 835))
    draw.text((1260, 765), "对目标函数构造可计算估计", font=small_font, fill="#475569")
    arrow(draw, (1200, 1075), (1200, 1195))

    # 在线策略的关键是参数更新会改变下一轮的状态分布。
    draw.line([(2150, 1315), (2290, 1315), (2290, 285), (2150, 285)], fill="#2563EB", width=8)
    draw.polygon([(2150, 285), (2182, 267), (2182, 303)], fill="#2563EB")
    draw.text((1685, 1115), "在线闭环：更新后的学生重新生成状态", font=small_font, fill="#1D4ED8")

    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(220, 220))


def add_picture_before(doc, anchor, image_path: Path, width_inches: float, ppr_template):
    paragraph = doc.add_paragraph()
    replace_ppr(paragraph, ppr_template)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    for doc_pr in paragraph._p.findall(".//" + qn("wp:docPr")):
        doc_pr.set("name", "三个研究问题关系图")
        doc_pr.set("descr", "问题一建立在线策略蒸馏基线，问题二和问题三分别研究散度方向与监督信息粒度，并共同形成统一实验矩阵和评测。")
    anchor._p.addprevious(paragraph._p)
    return paragraph


def apply_old_citation_mapping(doc: Document, references_heading) -> None:
    mapping = {
        1: 1,
        2: 2,
        3: 3,
        4: 4,
        5: None,
        6: 5,
        7: 6,
        8: 7,
        9: 8,
        10: None,
        11: 10,
        12: None,
        13: None,
        14: 11,
        15: None,
        16: 12,
        17: 13,
        18: 14,
        19: 15,
        20: 16,
        21: 17,
        22: None,
        23: 18,
        24: None,
        25: 19,
        26: 20,
        27: 21,
    }

    def repl(match):
        value = mapping.get(int(match.group(1)))
        return "" if value is None else f"[{value}]"

    for paragraph in doc.paragraphs:
        if paragraph._p is references_heading._p:
            break
        if paragraph.style and paragraph.style.name.lower().startswith("toc"):
            continue
        text = paragraph.text
        if not text or "[" not in text:
            continue
        replaced = re.sub(r"\[(\d+)\]", repl, text)
        replaced = re.sub(r" {2,}", " ", replaced)
        if replaced != text:
            set_paragraph_text(paragraph, replaced)


def build_static_toc(doc: Document, toc_pages=None) -> None:
    toc_title = find_paragraph(doc, "目 录")
    first_heading = find_paragraph(doc, "1．课题来源及研究的目的和意义")
    toc_paragraphs = [
        p for p in doc.paragraphs if p.style and p.style.name.lower().startswith("toc")
    ]
    if not toc_paragraphs:
        raise ValueError("未找到目录样式段落")
    toc1_ppr = deepcopy(next(p._p.pPr for p in toc_paragraphs if p.style.name.lower() == "toc 1"))
    toc2_ppr = deepcopy(next(p._p.pPr for p in toc_paragraphs if p.style.name.lower() == "toc 2"))
    remove_between(toc_title, first_heading)

    headings = []
    started = False
    for paragraph in doc.paragraphs:
        if paragraph._p is first_heading._p:
            started = True
        if not started:
            continue
        style = paragraph.style.name if paragraph.style else ""
        if style in {"Heading 1", "Heading 2"}:
            headings.append((paragraph.text, 1 if style == "Heading 1" else 2))

    toc_pages = toc_pages or {}
    for text, level in headings:
        paragraph = doc.add_paragraph()
        replace_ppr(paragraph, toc1_ppr if level == 1 else toc2_ppr)
        paragraph.add_run(text)
        paragraph.add_run("\t")
        paragraph.add_run(str(toc_pages.get(norm(text), 0)))
        first_heading._p.addprevious(paragraph._p)


def revise_report(report_path: Path, diagram_path: Path) -> None:
    doc = Document(report_path)
    references_heading = find_paragraph(doc, "8．主要参考文献")
    apply_old_citation_mapping(doc, references_heading)

    # 保存原始格式模板和需要继续使用的复杂对象。
    body_template = find_paragraph(doc, "本课题面向大语言模型后训练与工具增强应用中的实际需求，研究在线策略蒸馏（On-Policy Distillation，OPD）方法。研究将在学生模型自主生成的状态上引入教师模型的密集反馈，比较不同 KL 散度方向和监督信息粒度对训练效果、稳定性及资源开销的影响，并通过可交互的测试系统检验模型的实际能力。")
    list_template = find_paragraph(doc, "1. 梳理离线蒸馏、在线策略蒸馏、KL 散度和工具增强模型的发展脉络，明确 OPD 在 Function Calling 场景中的研究问题。")
    caption_template = find_paragraph(doc, "表 2-1 后训练方法的状态与监督来源对比")
    picture_template = find_paragraph(doc, "图 1 研究总体技术路线")._p.getprevious()
    body_ppr = deepcopy(body_template._p.pPr)
    list_ppr = deepcopy(list_template._p.pPr)
    caption_ppr = deepcopy(caption_template._p.pPr)
    picture_ppr = deepcopy(picture_template.find(qn("w:pPr")))

    table_2_1 = doc.tables[0]._tbl
    table_4_2 = doc.tables[2]._tbl
    eq_fkl = doc.paragraphs[102]._p
    eq_rkl = doc.paragraphs[103]._p
    eq_jsd = doc.paragraphs[106]._p
    eq_distill = doc.paragraphs[175]._p
    for element in [table_2_1, table_4_2, eq_fkl, eq_rkl, eq_jsd, eq_distill]:
        detach(element)

    # 第一部分：移出原 1.3/1.4，并将研究目的和意义前移为 1.3。
    h1_2 = find_paragraph(doc, "1.2 研究背景")
    h1_3_old = find_paragraph(doc, "1.3 传统后训练方法的局限")
    h1_4_old = find_paragraph(doc, "1.4 在线策略蒸馏的研究价值")
    h1_5 = find_paragraph(doc, "1.5 研究目的和意义")
    h2_main = find_paragraph(doc, "2．国内外在该方向的研究现状及分析")
    remove_between(h1_3_old, h1_4_old)
    remove_between(h1_4_old, h1_5)
    detach(h1_3_old._p)
    detach(h1_4_old._p)
    set_paragraph_text(h1_5, "1.3 研究目的和意义")

    remove_between(h1_2, h1_5)
    for text in [
        "大语言模型通过预训练获得语言知识，再通过监督微调、偏好优化、强化学习或知识蒸馏形成指令遵循和工具使用能力。知识蒸馏以教师模型的软分布补充硬标签，能够传递候选 token 之间的相对关系；序列级蒸馏则把教师生成序列作为训练目标[1][2]。",
        "自回归模型在训练时通常读取参考前缀，部署时却依赖自身历史输出，早期错误会改变后续状态。DAgger 和 ImitKD 表明，在学生实际到访的状态上查询专家或教师，可以减轻训练分布与生成分布之间的偏移[3][4]。在线策略蒸馏据此由学生生成轨迹，再由冻结教师在学生前缀上提供 token 级反馈。",
        "Function Calling 要求模型在给定工具 schema 下完成函数选择和参数生成，输出具有明确的结构约束，也可以通过解析、执行和端到端任务结果进行自动评测。API-Bank、APIGen、ToolACE 与 BFCL 分别从工具学习、可验证数据生成和多场景评测等方面建立了研究基础[18][19][20][21]。因此，该任务适合检验不同蒸馏方案在正确性、稳定性和资源成本上的差异。",
    ]:
        insert_paragraph_before(doc, h1_5, text, body_ppr)

    remove_between(h1_5, h2_main)
    for text in [
        "本课题以 Qwen3-0.6B 为学生模型、冻结的 Qwen3-4B 为教师模型，采用 LoRA 构建在线策略蒸馏流程，并在公开 Function Calling 数据上开展受控比较。",
        "1. 建立学生 rollout、教师评分、LoRA 更新和统一评测组成的在线策略蒸馏基线，明确状态分布、token 对齐和训练预算。",
        "2. 比较 Forward KL、Reverse KL 与代表性广义散度，分析散度方向对工具选择、参数生成和训练稳定性的影响。",
        "3. 比较全词表、Top-K 与 K1 sampled-token 监督，分析教师信息量、估计方差、任务效果和资源开销之间的权衡。",
        "研究的理论意义在于把学生状态分布、散度方向和监督信息粒度置于同一实验框架中考察；应用意义在于验证小模型通过参数高效后训练获得工具调用能力的可行性，为低成本部署提供可复现的比较依据。",
    ]:
        insert_paragraph_before(doc, h2_main, text, list_ppr if re.match(r"\d+\. ", text) else body_ppr)

    # 第二部分：压缩重复叙述，在末尾集中总结问题并引出研究问题。
    h2_1 = find_paragraph(doc, "2.1 大语言模型后训练和知识蒸馏")
    h2_2 = find_paragraph(doc, "2.2 分布偏移与在线策略蒸馏")
    h2_3 = find_paragraph(doc, "2.3 KL 散度方向与广义散度")
    h2_4 = find_paragraph(doc, "2.4 监督信息粒度与资源效率")
    h2_5 = find_paragraph(doc, "2.5 Function Calling 数据与评测研究")
    h3_main = find_paragraph(doc, "3．主要研究内容")
    for start, end in [(h2_1, h2_2), (h2_2, h2_3), (h2_3, h2_4), (h2_4, h2_5), (h2_5, h3_main)]:
        remove_between(start, end)

    for text in [
        "大语言模型后训练可按监督来源分为监督微调、知识蒸馏、偏好优化和强化学习。监督微调使用目标序列，知识蒸馏匹配教师分布，偏好优化使用回答偏好，强化学习依据任务奖励更新策略。本文关注知识蒸馏，不把偏好优化和强化学习列为主要实验变量。",
        "经典知识蒸馏通过温度缩放保留教师对多个候选的相对偏好[1]，序列级蒸馏把教师生成序列用于训练学生[2]，f-divergence 视角进一步统一了多种分布匹配目标[5]。这些工作奠定了生成模型蒸馏的基础，但固定语料或教师轨迹不能充分覆盖学生部署时访问的状态。",
    ]:
        insert_paragraph_before(doc, h2_2, text, body_ppr)
    insert_paragraph_before(doc, h2_2, "表 2-1 后训练方法的状态与监督来源对比", caption_ppr)
    insert_element_before(h2_2, table_2_1)
    insert_paragraph_before(
        doc,
        h2_2,
        "表 2-1 表明，OPD 与离线蒸馏的关键差别是训练状态由学生策略生成，教师在这些状态上提供密集反馈。本文以 LoRA-OPD 为主线，并将其他后训练方法作为概念对照。",
        body_ppr,
    )

    for text in [
        "自回归生成会把模型先前输出重新作为后续输入，因此训练前缀与部署前缀的差异会逐步累积。DAgger 在学习器到访状态上查询专家动作[3]，ImitKD 将该思想用于自回归知识蒸馏[4]，为在线策略蒸馏提供了直接的方法基础。",
        "GKD 在学生自生成轨迹上比较教师与学生分布，并将散度目标作为可配置变量[6]；MiniLLM 从序列级 Reverse KL 出发强调学生策略访问区域[7]；DistiLLM 通过 skew KL 与样本复用讨论效果和成本的平衡[8]。",
        "DistiLLM-2 进一步用对比式目标联合利用教师与学生生成数据[9]。近期研究还指出，师生推理模式、提示覆盖、轨迹长度和高概率 token 重叠都会影响 OPD 的有效性[10]。这说明学生状态、损失目标和教师信息预算需要在同一协议下分别考察。",
        "现有工作往往同时改变模型规模、数据来源、采样方法和损失函数，论文间的结果难以直接横向比较。本文据此固定师生模型、Function Calling 数据、训练预算和解码协议，只改变散度方向与监督信息粒度。",
    ]:
        insert_paragraph_before(doc, h2_3, text, body_ppr)

    insert_paragraph_before(
        doc,
        h2_4,
        "在同一生成前缀上，教师与学生分别给出条件 token 分布。Forward KL 与 Reverse KL 定义为[13]：",
        body_ppr,
    )
    insert_element_before(h2_4, eq_fkl)
    insert_element_before(h2_4, eq_rkl)
    insert_paragraph_before(
        doc,
        h2_4,
        "Forward KL 以教师概率加权，通常更强调覆盖教师分布；Reverse KL 以学生概率加权，更关注学生当前高概率区域。mode-covering 与 mode-seeking 只能作为行为倾向的解释，实际结果还受模型容量、温度、采样策略和支持集影响[6][7]。",
        body_ppr,
    )
    insert_paragraph_before(
        doc,
        h2_4,
        "Jensen-Shannon 散度通过教师分布、学生分布及其混合分布定义[14]：",
        body_ppr,
    )
    insert_element_before(h2_4, eq_jsd)
    insert_paragraph_before(
        doc,
        h2_4,
        "广义散度允许在覆盖与集中之间调节。GKD、DistiLLM 与 HPD 分别从可调散度、skew KL 和混合 KL 方向研究这一问题[6][8][11]。本文将散度定义、支持集和数值稳定化方式一并记录，避免把实现差异误判为目标函数差异。",
        body_ppr,
    )

    for text in [
        "全词表监督保留教师在每个有效位置上的完整分布，信息最充分，但需要较高的计算、显存与传输开销。Top-K 只保留高概率支持集，需要明确 K 值、支持集来源、截断后的重归一化和尾部概率处理。K1 sampled-token 用单个采样 token 估计期望，开销最低，但估计方差与 proposal 分布会直接影响稳定性[10]。",
        "LoRA 通过低秩增量减少学生侧可训练参数[15]，使受控对照实验能够在有限算力下实施。它只改变参数更新方式，不替代蒸馏目标；教师前向、学生 rollout 和损失计算仍是主要成本。",
        "监督信息越稀疏，并不必然意味着总体效率越高。Top-K 或 K1 如果造成收敛变慢、训练波动或任务性能下降，节省的单步开销可能被更多训练步数抵消。因此，研究需要同时报告任务指标、有效 token 吞吐、教师调用量、峰值显存和训练时间。",
    ]:
        insert_paragraph_before(doc, h2_5, text, body_ppr)

    for text in [
        "ReAct 把语言模型的推理与外部行动交替组织，为工具增强模型提供了通用范式[17]。API-Bank 构建工具规划、检索与调用基准[18]；APIGen 和 ToolACE 分别强调可执行验证与多层数据检查[19][20]。这些研究说明，Function Calling 数据不能只按文本样本处理，还需要核验 schema、参数类型、调用可执行性和对话模板。",
        "BFCL 覆盖单次、并行、串行和拒答等调用场景，并通过结构和语义规则评估函数与参数[21]。总体准确率仍可能掩盖错误来源，因此本文将分别统计格式合法率、函数选择准确率、参数正确率、执行成功率和端到端任务成功率。",
        "训练集与测试集还应按工具 schema、问题模板和数据来源进行隔离，防止相似 API 或合成模板造成泄漏。数据准备阶段将保存来源、版本、许可证、字段映射、去重规则和版本指纹，使方法比较能够追溯。",
    ]:
        insert_paragraph_before(doc, h3_main, text, body_ppr)

    set_paragraph_text(h1_3_old, "2.6 现有方法存在的问题")
    set_paragraph_text(h1_4_old, "2.7 本课题研究问题")
    insert_element_before(h3_main, h1_3_old._p)
    for text in [
        "综合上述研究，现有方法仍存在三方面不足：",
        "1. 状态分布与实验协议缺少统一控制。不同研究常同时改变学生采样、数据来源、模型规模和训练预算，难以判断收益究竟来自 on-policy 状态还是其他条件。",
        "2. 散度方向与监督信息粒度相互耦合。Top-K 截断和 K1 采样会改变实际优化对象及梯度方差，如果不明确支持集和估计方式，损失名称相同也可能得到不同结果。",
        "3. token 级分布匹配与端到端工具调用之间缺少系统联系。较低的蒸馏损失不一定带来更高的函数、参数或执行正确率，训练成本与部署性能也可能呈现不同趋势。",
    ]:
        insert_paragraph_before(doc, h3_main, text, list_ppr if re.match(r"\d+\. ", text) else body_ppr)
    insert_element_before(h3_main, h1_4_old._p)
    for text in [
        "针对上述不足，本课题归纳出三个相互衔接的研究问题：",
        "1. 如何在学生自身生成的 Function Calling 状态上建立接口统一、变量可控的在线策略蒸馏基线？",
        "2. 在相同学生状态与训练预算下，Forward KL、Reverse KL 和广义散度会产生怎样的学习行为与任务差异？",
        "3. 全词表、Top-K 和 K1 sampled-token 如何在任务效果、训练稳定性与教师反馈开销之间取得平衡？",
        "这三个问题分别对应论文第 2、3、4 章，并在第 3 部分进一步说明研究内容及其关系。",
    ]:
        insert_paragraph_before(doc, h3_main, text, list_ppr if re.match(r"\d+\. ", text) else body_ppr)

    # 第三部分：仅保留三个问题，并加入关系图。
    h3_1 = find_paragraph(doc, "3.1 统一的在线策略蒸馏研究流程")
    h3_2 = find_paragraph(doc, "3.2 三类散度目标")
    h3_3 = find_paragraph(doc, "3.3 三种监督信息粒度/估计方式")
    h4_main = find_paragraph(doc, "4．研究方案")
    for heading in [h3_1, h3_2, h3_3]:
        detach(heading._p)
    remove_between(h3_main, h4_main)

    insert_paragraph_before(
        doc,
        h4_main,
        "研究内容围绕三个问题组织。三者分别对应论文第 2、3、4 章：先建立可复现的在线策略蒸馏基线，再比较散度方向，最后分析监督信息粒度下的效果与效率权衡。",
        body_ppr,
    )
    set_paragraph_text(h3_1, "3.1 学生状态分布下的在线蒸馏基线问题")
    insert_element_before(h4_main, h3_1._p)
    for text in [
        "该问题对应论文第 2 章。研究以学生自主生成的 Function Calling 轨迹为训练状态，冻结教师在相同前缀上提供 token 分布，学生仅更新 LoRA 参数。核心是统一数据格式、消息模板、token 对齐、rollout 规则、有效位置 mask 和训练预算。",
        "实验先建立 SFT 或离线蒸馏参照，再验证学生 rollout、教师评分和梯度更新能否稳定闭环。输出包括可复现的训练流程、状态分布统计和基础任务指标，为后续两个问题提供共同基线。",
    ]:
        insert_paragraph_before(doc, h4_main, text, body_ppr)

    set_paragraph_text(h3_2, "3.2 散度方向与学习行为差异问题")
    insert_element_before(h4_main, h3_2._p)
    for text in [
        "该问题对应论文第 3 章。在固定学生状态、教师模型和监督信息范围的条件下，比较 Forward KL、Reverse KL 和代表性广义散度，分析不同目标对概率覆盖、模式集中、熵变化和梯度稳定性的影响。",
        "结果不仅报告总体任务得分，还按函数选择、参数生成、格式合法性和拒答行为分类，并结合师生分布差异解释性能变化，避免只依据损失数值判断方法优劣。",
    ]:
        insert_paragraph_before(doc, h4_main, text, body_ppr)

    set_paragraph_text(h3_3, "3.3 监督信息粒度与效果效率权衡问题")
    insert_element_before(h4_main, h3_3._p)
    for text in [
        "该问题对应论文第 4 章。在相同散度定义和训练预算下，比较全词表、Top-K 与 K1 sampled-token。研究重点是支持集截断、尾部概率处理和单样本估计如何改变有效监督、梯度方差及教师计算开销。",
        "评价同时覆盖 Function Calling 正确性、训练稳定性、教师反馈量、训练吞吐、峰值显存和最终推理性能，从而判断稀疏监督节省的资源是否足以补偿可能的性能损失。",
        "三个问题的关系如图 1 所示。问题一界定教师反馈所作用的学生状态分布；在该分布上，问题二决定师生分布如何匹配，问题三决定匹配目标如何由有限教师信息估计。三者共同影响更新梯度、优化稳定性、Function Calling 正确性与教师反馈成本，更新后的学生又会改变下一轮状态分布，形成在线闭环。",
    ]:
        insert_paragraph_before(doc, h4_main, text, body_ppr)

    create_relationship_diagram(diagram_path)
    diagram_paragraph = add_picture_before(doc, h4_main, diagram_path, 6.15, picture_ppr)
    diagram_caption = insert_paragraph_before(doc, h4_main, "图 1 三个研究问题及其关系", caption_ppr)
    set_keep_with_next(diagram_paragraph, True)
    set_keep_with_next(diagram_caption, False)

    # 第四部分：在总体技术路线与监督实现中增加形式化描述。
    h4_1 = find_paragraph(doc, "4.1 总体技术路线")
    h4_2 = find_paragraph(doc, "4.2 模型和数据准备")
    route_caption = find_paragraph(doc, "图 1 研究总体技术路线")
    route_image = route_caption._p.getprevious()
    detach(route_image)
    remove_between(h4_1, h4_2)
    for text in [
        "研究总体技术路线如图 2 所示。给定输入与当前学生参数，学生首先按自身策略生成序列，并定义第 t 个位置的状态：",
    ]:
        insert_paragraph_before(doc, h4_2, text, body_ppr)
    insert_element_before(h4_2, rollout_equation())
    insert_paragraph_before(doc, h4_2, "冻结教师参数记为 φ。教师与学生在相同状态上给出的 token 分布分别为：", body_ppr)
    insert_element_before(h4_2, distribution_equation())
    insert_paragraph_before(doc, h4_2, "统一的在线策略蒸馏目标写为：", body_ppr)
    insert_element_before(h4_2, objective_equation())
    for text in [
        "其中，𝒟 表示训练输入分布；有效位置掩码由消息边界、工具边界和响应位置共同确定；散度参数指定 Forward KL、Reverse KL 或广义散度。该形式把学生状态分布、散度方向和监督信息粒度统一到同一目标中。",
        "每项实验同时记录研究定义、数据定义、模型定义和运行定义，包括散度及支持集、数据版本与划分、师生模型及模板、随机种子与计算环境，确保结果能够追溯。",
    ]:
        insert_paragraph_before(doc, h4_2, text, body_ppr)
    insert_element_before(h4_2, route_image)
    insert_paragraph_before(doc, h4_2, "图 2 研究总体技术路线", caption_ppr)

    p = find_paragraph(doc, "在线策略蒸馏的训练、反馈与评测关系如图 2 所示。")
    set_paragraph_text(p, "在线策略蒸馏的训练、反馈与评测关系如图 3 所示。")
    set_paragraph_text(find_paragraph(doc, "图 2 在线策略蒸馏训练与性能评测流程"), "图 3 在线策略蒸馏训练与性能评测流程")

    h4_4 = find_paragraph(doc, "4.4 损失与监督信息粒度/估计方式的定义原则")
    h4_5 = find_paragraph(doc, "4.5 评价指标和公平性")
    remove_between(h4_4, h4_5)
    insert_paragraph_before(
        doc,
        h4_5,
        "在全词表方案中，教师与学生在每个有效位置上使用完整词表分布，批次级目标按有效 token 加权平均：",
        body_ppr,
    )
    insert_element_before(h4_5, eq_distill)
    insert_paragraph_before(
        doc,
        h4_5,
        "Top-K 方案先定义教师高概率支持集及其概率质量：",
        body_ppr,
    )
    insert_element_before(h4_5, topk_equation())
    insert_paragraph_before(doc, h4_5, "在支持集内重归一化的教师分布为：", body_ppr)
    insert_element_before(h4_5, topk_normalized_equation())
    insert_paragraph_before(
        doc,
        h4_5,
        "K1 sampled-token 使用 proposal 分布采样一个 token。对教师分布下某函数的期望，其通用重要性采样估计写为：",
        body_ppr,
    )
    insert_element_before(h4_5, k1_equation())
    insert_paragraph_before(
        doc,
        h4_5,
        "该估计要求 proposal 分布覆盖教师非零支持集，并记录采样分布、概率比裁剪、随机种子和估计方差。Top-K 与 K1 不仅减少信息量，也可能改变优化目标或梯度噪声，因此必须与全词表方案使用相同的数据、rollout 条件和总训练预算。",
        body_ppr,
    )
    insert_paragraph_before(doc, h4_5, "表 4-2 散度目标与监督信息粒度比较矩阵", caption_ppr)
    insert_element_before(h4_5, table_4_2)
    insert_paragraph_before(
        doc,
        h4_5,
        "表 4-2 构成主要实验矩阵。每个组合都记录具体散度公式、K 值、支持集来源、重归一化方式、K1 proposal、mask、数值精度和梯度处理；正式训练前通过 tokenizer、模板回环、损失数值和梯度传播测试验证实现。",
        body_ppr,
    )

    schedule_intro = find_paragraph(doc, "本课题自 2026 年 9 月开题申请起，至 2027 年 5 月结题申请与答辩结束，按实际培养环节组织实施。开题申请时间为 2026 年 9 月 14 日至 9 月 30 日，中期申请预计在 2027 年 3 月，结题申请与答辩预计在 2027 年 5 月。进度安排如图 3 和表 5-1 所示。")
    set_paragraph_text(
        schedule_intro,
        "本课题自 2026 年 9 月开题申请起，至 2027 年 5 月结题申请与答辩结束，按实际培养环节组织实施。开题申请时间为 2026 年 9 月 14 日至 9 月 30 日，中期申请预计在 2027 年 3 月，结题申请与答辩预计在 2027 年 5 月。进度安排如图 4 和表 5-1 所示。",
    )
    set_paragraph_text(find_paragraph(doc, "图 3 项目总体进度安排"), "图 4 项目总体进度安排")

    data_preparation = find_paragraph(
        doc,
        "数据方面，核验公开 Function Calling 数据的来源、许可证、字段含义和版本，再把实际字段转换为统一的用户消息、工具 schema、调用响应与工具结果表示。训练集、验证集和测试集按照来源与模板隔离，并记录划分比例、子集范围和去重规则。API-Bank、ToolLLM/ToolBench、APIGen 和 ToolACE 为数据选择与构造提供参考[18][19][20]。",
    )
    set_paragraph_text(
        data_preparation,
        "数据方面，核验公开 Function Calling 数据的来源、许可证、字段含义和版本，再把实际字段转换为统一的用户消息、工具 schema、调用响应与工具结果表示。训练集、验证集和测试集按照来源与模板隔离，并记录划分比例、子集范围和去重规则。API-Bank、APIGen 和 ToolACE 为数据选择与构造提供参考[18][19][20]。",
    )
    set_plain_paragraph_text(
        find_paragraph(doc, "表 4-3 评价指标体系"),
        "表 4-3 评价指标体系",
    )
    feasibility = find_paragraph(
        doc,
        "Qwen3-0.6B 学生、Qwen3-4B 教师和 LoRA 后训练构成了可实施的模型组合，PyTorch、Transformers、PEFT 与 vLLM 能够提供训练和部署所需的基础工具。研究将先进行小规模验证，在模型、数据、接口和评测规则固定后逐步扩大训练规模，因而具备清晰的实施路径和验证条件。",
    )
    set_paragraph_text(
        feasibility,
        "Qwen3-0.6B 学生、Qwen3-4B 教师和 LoRA 后训练构成了可实施的模型组合，PyTorch、Transformers、PEFT 与 vLLM 能够提供训练和部署所需的基础工具[16]。研究将先进行小规模验证，在模型、数据、接口和评测规则固定后逐步扩大训练规模，因而具备清晰的实施路径和验证条件。",
    )

    # 精选参考文献：删除技术博客和与本报告主线关系较弱的条目。
    remove_between(references_heading, None)
    sect_pr = doc.element.body.sectPr
    for index, entry in enumerate(REFERENCE_ENTRIES, start=1):
        paragraph = doc.add_paragraph()
        replace_ppr(paragraph, list_ppr)
        paragraph.add_run(f"[{index}] {entry}")
        sect_pr.addprevious(paragraph._p)

    # 让 Word 打开文档时主动更新域；目录同时写入可视的静态缓存以便无 Word 环境查看。
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    build_static_toc(doc)

    doc.core_properties.modified = __import__("datetime").datetime.now()
    temp_path = report_path.with_suffix(".tmp.docx")
    doc.save(temp_path)
    os.replace(temp_path, report_path)


def update_toc_pages(report_path: Path, page_map: dict[str, int]) -> None:
    doc = Document(report_path)
    build_static_toc(doc, page_map)
    temp_path = report_path.with_suffix(".tmp.docx")
    doc.save(temp_path)
    os.replace(temp_path, report_path)


def prefix_reference_numbers(report_path: Path) -> None:
    doc = Document(report_path)
    heading = find_paragraph(doc, "8．主要参考文献")
    started = False
    index = 0
    for paragraph in doc.paragraphs:
        if paragraph._p is heading._p:
            started = True
            continue
        if not started or not paragraph.text.strip():
            continue
        index += 1
        text = re.sub(r"^\[\d+\]\s*", "", paragraph.text.strip())
        set_paragraph_text(paragraph, f"[{index}] {text}")
    if index != len(REFERENCE_ENTRIES):
        raise ValueError(f"参考文献数量异常：{index}")
    temp_path = report_path.with_suffix(".tmp.docx")
    doc.save(temp_path)
    os.replace(temp_path, report_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("report", nargs="?", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--diagram", type=Path, default=Path("/tmp/opd_research_questions.png"))
    parser.add_argument("--toc-pages", type=Path)
    parser.add_argument("--fix-reference-numbers", action="store_true")
    args = parser.parse_args()

    if args.fix_reference_numbers:
        prefix_reference_numbers(args.report)
    elif args.toc_pages:
        import json

        update_toc_pages(args.report, json.loads(args.toc_pages.read_text(encoding="utf-8")))
    else:
        revise_report(args.report, args.diagram)


if __name__ == "__main__":
    main()
